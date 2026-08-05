#!/usr/bin/env python
"""Execution Verification Script for Report Source Integration against PostgreSQL 17.

Tests PL.01, PL.02, PL.03, and Combined source export logic using real PostgreSQL
test database 'cangvu_ui_test', generating XLSX samples and execution logs.
"""
import os
import sys
import json
import io
import hashlib
import secrets
import subprocess
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

# Resolve project root
CURRENT_FILE = Path(__file__).resolve()
EVIDENCE_DIR = CURRENT_FILE.parent
PROJECT_ROOT = CURRENT_FILE.parents[4]
sys.path.insert(0, str(PROJECT_ROOT))

# Resolve DATABASE_URL from environment variable
TEST_DB_URL = os.getenv("QLXL_UI_TEST_DATABASE_URL")
if not TEST_DB_URL:
    print("[ERROR] QLXL_UI_TEST_DATABASE_URL environment variable is required for verification.", file=sys.stderr)
    sys.exit(1)
os.environ["DATABASE_URL"] = TEST_DB_URL
os.environ["SECRET_KEY"] = os.getenv("SECRET_KEY") or secrets.token_urlsafe(32)

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from backend.app import app
from backend.auth import create_access_token
from backend.database import SessionLocal
from backend.models import User, ReportingUnitUser
from sqlalchemy import text


def sha256_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def current_git_sha() -> str:
    result = subprocess.run(
        ["git", "rev-parse", "HEAD"],
        cwd=PROJECT_ROOT,
        check=True,
        capture_output=True,
        text=True,
    )
    return result.stdout.strip()


def worktree_is_dirty() -> bool:
    result = subprocess.run(
        ["git", "status", "--porcelain"],
        cwd=PROJECT_ROOT,
        check=True,
        capture_output=True,
        text=True,
    )
    return bool(result.stdout.strip())

def main():
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass
    log_lines = []
    def log(msg: str):
        print(msg)
        log_lines.append(msg)

    log("======================================================================")
    execution_time = datetime.now(ZoneInfo("Asia/Ho_Chi_Minh")).isoformat(timespec="seconds")
    git_sha = current_git_sha()
    worktree_dirty = worktree_is_dirty()
    log("REPORT SOURCE EXPORT ENGINE VERIFICATION (PostgreSQL)")
    log(f"Execution Time: {execution_time}")
    log(f"Git SHA: {git_sha}")
    log(f"Worktree dirty: {str(worktree_dirty).lower()}")
    log(f"Command: python {CURRENT_FILE.relative_to(PROJECT_ROOT).as_posix()}")
    log(f"Target Database: cangvu_ui_test")
    log("======================================================================\n")

    client = TestClient(app)

    db = SessionLocal()
    try:
        server_version = db.execute(text("SHOW server_version")).scalar_one()
        admin_user = db.query(User).filter_by(username="admin").first()
        if not admin_user:
            raise RuntimeError("Admin user not found in cangvu_ui_test. Run setup_ui_test_env.py first.")
        ru_user = db.query(ReportingUnitUser).filter_by(user_id=admin_user.id).first()
        reporting_unit_id = str(ru_user.reporting_unit_id) if ru_user else "1"
        token = create_access_token({
            "sub": admin_user.username,
            "role": admin_user.role,
            "org_id": admin_user.organization_id,
            "pwd_ts": admin_user.password_changed_at,
        })
    finally:
        db.close()

    headers = {
        "Authorization": f"Bearer {token}",
        "X-Reporting-Unit-ID": reporting_unit_id,
    }

    log(f"PostgreSQL server_version: {server_version}\n")
    verification_results = {
        "metadata": {
            "execution_time": execution_time,
            "git_sha": git_sha,
            "worktree_dirty": worktree_dirty,
            "postgresql_server_version": server_version,
            "database": "cangvu_ui_test",
            "source_sha256": {
                "backend/reports_api.py": sha256_file(PROJECT_ROOT / "backend" / "reports_api.py"),
                "backend/historical_api.py": sha256_file(PROJECT_ROOT / "backend" / "historical_api.py"),
                "frontend/app.js": sha256_file(PROJECT_ROOT / "frontend" / "app.js"),
                "frontend/index.html": sha256_file(PROJECT_ROOT / "frontend" / "index.html"),
            },
        }
    }

    # ------------------------------------------------------------------
    # TEST 1: PL.01 Export Logic & Contract Enforcement
    # ------------------------------------------------------------------
    log("--- TEST 1: PL.01 (LIVE Only Contract) ---")
    res1_live = client.get("/api/reports/appendix1?source=live&from=2026-05-01&to=2026-05-31", headers=headers)
    log(f"1a. GET /api/reports/appendix1?source=live -> Status {res1_live.status_code}")
    assert res1_live.status_code == 200, f"Expected 200, got {res1_live.status_code}: {res1_live.text}"
    pl01_path = EVIDENCE_DIR / "PL01_LIVE_sample.xlsx"
    pl01_path.write_bytes(res1_live.content)
    log(f"    Saved sample: {pl01_path.name} ({len(res1_live.content)} bytes)")

    res1_invalid = client.get("/api/reports/appendix1?source=historical&from=2026-05-01&to=2026-05-31", headers=headers)
    log(f"1b. GET /api/reports/appendix1?source=historical -> Status {res1_invalid.status_code}")
    assert res1_invalid.status_code == 422, f"Expected 422 refusal, got {res1_invalid.status_code}"
    log(f"    Refusal detail: {res1_invalid.json().get('detail')}")

    invalid_log_path = EVIDENCE_DIR / "422_pl01_invalid_source.log"
    invalid_log_path.write_text(f"HTTP GET /api/reports/appendix1?source=historical\nStatus: {res1_invalid.status_code}\nResponse: {res1_invalid.text}\n", encoding="utf-8")

    verification_results["PL01_LIVE"] = {
        "status": "PASS",
        "live_export_status": res1_live.status_code,
        "live_workbook_sha256": sha256_bytes(res1_live.content),
        "invalid_source_refusal_status": res1_invalid.status_code,
        "refusal_message": res1_invalid.json().get("detail")
    }

    # ------------------------------------------------------------------
    # TEST 2: PL.02 Historical TOS Export (April 2026)
    # ------------------------------------------------------------------
    log("\n--- TEST 2: PL.02 Historical TOS Export (2026-04) ---")
    res2_tos = client.get("/api/reports/appendix2?source=historical&to=2026-04-30", headers=headers)
    log(f"2a. GET /api/reports/appendix2?source=historical&to=2026-04-30 -> Status {res2_tos.status_code}")
    assert res2_tos.status_code == 200, f"Expected 200, got {res2_tos.status_code}: {res2_tos.text}"
    pl02_path = EVIDENCE_DIR / "PL02_TOS_sample.xlsx"
    pl02_path.write_bytes(res2_tos.content)
    log(f"    Saved sample: {pl02_path.name} ({len(res2_tos.content)} bytes)")

    # Inspect openpyxl sheet to confirm historical metrics
    wb2 = load_workbook(io.BytesIO(res2_tos.content), data_only=True)
    sheet2 = wb2.active

    # Find data row containing '- Cảng Tân Thuận' or port name
    target_row = None
    for r in range(1, sheet2.max_row + 1):
        v = sheet2.cell(row=r, column=2).value
        if v and "- Cảng" in str(v):
            target_row = r
            break

    assert target_row is not None, "Could not find port row in PL.02 Excel"
    vessel_cell = sheet2.cell(row=target_row, column=2).value  # Port / Facility Name
    container_tons_cell = sheet2.cell(row=target_row, column=3).value  # Col 3: container_tons
    container_teu_cell = sheet2.cell(row=target_row, column=4).value   # Col 4: container_teu
    dry_tons_cell = sheet2.cell(row=target_row, column=7).value        # Col 7: dry_tons (Unsupported historical)
    liquid_tons_cell = sheet2.cell(row=target_row, column=9).value     # Col 9: liquid_tons (Unsupported historical)

    log(f"    Excel Row {target_row} Inspection:")
    log(f"      - Facility Name: '{vessel_cell}'")
    log(f"      - Container Tonnes (TOS): {container_tons_cell}")
    log(f"      - Container TEU (TOS): {container_teu_cell}")
    log(f"      - Unsupported Dry Tonnes cell state: {dry_tons_cell} (Must be None/Blank, not 0)")
    log(f"      - Unsupported Liquid Tonnes cell state: {liquid_tons_cell} (Must be None/Blank, not 0)")

    assert dry_tons_cell is None, f"Expected None/Blank for unsupported dry_tons metric, got {dry_tons_cell}"
    assert liquid_tons_cell is None, f"Expected None/Blank for unsupported liquid_tons metric, got {liquid_tons_cell}"

    verification_results["PL02_TOS"] = {
        "status": "PASS",
        "export_status": res2_tos.status_code,
        "workbook_sha256": sha256_bytes(res2_tos.content),
        "container_tons": container_tons_cell,
        "container_teu": container_teu_cell,
        "unsupported_dry_tons_is_blank": dry_tons_cell is None,
        "unsupported_liquid_tons_is_blank": liquid_tons_cell is None
    }

    # ------------------------------------------------------------------
    # TEST 3: PL.03 Historical TOS Export & Per-Call Row Integrity (July 2026)
    # ------------------------------------------------------------------
    log("\n--- TEST 3: PL.03 Historical TOS Export & Per-Call Per-Voyage Rows (2026-07) ---")
    res3_tos = client.get("/api/reports/appendix3?source=historical&from=2026-07-01&to=2026-07-31", headers=headers)
    log(f"3a. GET /api/reports/appendix3?source=historical&from=2026-07-01&to=2026-07-31 -> Status {res3_tos.status_code}")
    assert res3_tos.status_code == 200, f"Expected 200, got {res3_tos.status_code}: {res3_tos.text}"
    pl03_path = EVIDENCE_DIR / "PL03_TOS_sample.xlsx"
    pl03_path.write_bytes(res3_tos.content)
    log(f"    Saved sample: {pl03_path.name} ({len(res3_tos.content)} bytes)")

    wb3 = load_workbook(io.BytesIO(res3_tos.content), data_only=True)
    sheet3 = wb3.active

    # Collect rows for Salan SG-9999 (Registration SG-9999 in Column C / col 3)
    salan_rows = []
    for r in range(1, sheet3.max_row + 1):
        reg = sheet3.cell(row=r, column=3).value
        if reg and "SG-9999" in str(reg):
            atb = sheet3.cell(row=r, column=33).value # AG col ATB
            atd = sheet3.cell(row=r, column=34).value # AH col ATD
            tonnes = sheet3.cell(row=r, column=15).value # O col tonnes
            salan_rows.append((r, str(reg), str(atb or ""), str(atd or ""), tonnes))

    log(f"    Found {len(salan_rows)} separate voyage rows for 'SG-9999':")
    for r, reg, atb, atd, tonnes in salan_rows:
        log(f"      - Row {r}: Reg='{reg}', ATB='{atb}', ATD='{atd}', Tonnes={tonnes}")

    assert len(salan_rows) == 2, f"Expected 2 distinct voyage rows for SG-9999 in July 2026, got {len(salan_rows)}"
    row1_atb, row2_atb = salan_rows[0][2], salan_rows[1][2]
    assert row1_atb != row2_atb, f"ATB timestamps must be distinct between voyages ({row1_atb} vs {row2_atb})"

    verification_results["PL03_TOS"] = {
        "status": "PASS",
        "export_status": res3_tos.status_code,
        "workbook_sha256": sha256_bytes(res3_tos.content),
        "distinct_salan_voyage_rows": len(salan_rows),
        "voyage_1_atb": row1_atb,
        "voyage_2_atb": row2_atb,
        "voyages_not_merged": len(salan_rows) == 2 and row1_atb != row2_atb
    }

    # ------------------------------------------------------------------
    # TEST 4: COMBINED Source Export (Non-Overlapping & Overlapping)
    # ------------------------------------------------------------------
    log("\n--- TEST 4: COMBINED Mode Export (Non-Overlap Success & Overlap Refusal) ---")

    # 4a/4b. PL.02 non-overlap success and overlapping-month refusal.
    res4_pl02_non_overlap = client.get(
        "/api/reports/appendix2?source=combined&to=2026-05-31", headers=headers
    )
    log(
        "4a. GET /api/reports/appendix2?source=combined&to=2026-05-31 "
        f"-> Status {res4_pl02_non_overlap.status_code}"
    )
    assert res4_pl02_non_overlap.status_code == 200, (
        f"Expected PL.02 combined 200, got {res4_pl02_non_overlap.status_code}"
    )
    pl02_combined_path = EVIDENCE_DIR / "PL02_COMBINED_sample.xlsx"
    pl02_combined_path.write_bytes(res4_pl02_non_overlap.content)
    log(f"    Saved sample: {pl02_combined_path.name} ({len(res4_pl02_non_overlap.content)} bytes)")

    res4_pl02_overlap = client.get(
        "/api/reports/appendix2?source=combined&to=2026-07-31", headers=headers
    )
    log(
        "4b. GET /api/reports/appendix2?source=combined&to=2026-07-31 "
        f"-> Status {res4_pl02_overlap.status_code}"
    )
    assert res4_pl02_overlap.status_code == 409, (
        f"Expected PL.02 overlap 409, got {res4_pl02_overlap.status_code}"
    )
    pl02_overlap_log_path = EVIDENCE_DIR / "409_appendix2_overlap_block.log"
    pl02_overlap_log_path.write_text(
        "HTTP GET /api/reports/appendix2?source=combined&to=2026-07-31\n"
        f"Status: {res4_pl02_overlap.status_code}\nResponse: {res4_pl02_overlap.text}\n",
        encoding="utf-8",
    )

    # 4c/4d. PL.03 non-overlap success and overlapping-month refusal.
    res4_pl03_non_overlap = client.get(
        "/api/reports/appendix3?source=combined&from=2026-04-01&to=2026-05-31",
        headers=headers,
    )
    log(
        "4c. GET /api/reports/appendix3?source=combined&from=2026-04-01&to=2026-05-31 "
        f"-> Status {res4_pl03_non_overlap.status_code}"
    )
    assert res4_pl03_non_overlap.status_code == 200, (
        f"Expected PL.03 combined 200, got {res4_pl03_non_overlap.status_code}"
    )
    pl03_combined_path = EVIDENCE_DIR / "PL03_COMBINED_sample.xlsx"
    pl03_combined_path.write_bytes(res4_pl03_non_overlap.content)
    log(f"    Saved sample: {pl03_combined_path.name} ({len(res4_pl03_non_overlap.content)} bytes)")

    res4_pl03_overlap = client.get(
        "/api/reports/appendix3?source=combined&from=2026-07-01&to=2026-07-31",
        headers=headers,
    )
    log(
        "4d. GET /api/reports/appendix3?source=combined&from=2026-07-01&to=2026-07-31 "
        f"-> Status {res4_pl03_overlap.status_code}"
    )
    assert res4_pl03_overlap.status_code == 409, (
        f"Expected PL.03 overlap 409, got {res4_pl03_overlap.status_code}"
    )
    log(f"    PL.03 conflict message: {res4_pl03_overlap.json().get('detail')}")

    pl03_overlap_log_path = EVIDENCE_DIR / "409_appendix3_overlap_block.log"
    pl03_overlap_log_path.write_text(
        "HTTP GET /api/reports/appendix3?source=combined&from=2026-07-01&to=2026-07-31\n"
        f"Status: {res4_pl03_overlap.status_code}\nResponse: {res4_pl03_overlap.text}\n",
        encoding="utf-8",
    )

    verification_results["COMBINED_MODE"] = {
        "status": "PASS",
        "appendix2_non_overlapping_export_status": res4_pl02_non_overlap.status_code,
        "appendix2_workbook_sha256": sha256_bytes(res4_pl02_non_overlap.content),
        "appendix2_overlapping_refusal_status": res4_pl02_overlap.status_code,
        "appendix2_conflict_message": res4_pl02_overlap.json().get("detail"),
        "appendix3_non_overlapping_export_status": res4_pl03_non_overlap.status_code,
        "appendix3_workbook_sha256": sha256_bytes(res4_pl03_non_overlap.content),
        "appendix3_overlapping_refusal_status": res4_pl03_overlap.status_code,
        "appendix3_conflict_message": res4_pl03_overlap.json().get("detail"),
        "double_counting_prevented": (
            res4_pl02_overlap.status_code == 409 and res4_pl03_overlap.status_code == 409
        ),
    }

    # Save summary execution log & workbook verification receipt
    log("\n======================================================================")
    log(f"ALL POSTGRESQL {server_version} REPORT ENGINE EXECUTABLE TESTS PASSED SUCCESSFULLY")
    log("======================================================================\n")

    log_path = EVIDENCE_DIR / "test_execution_log.txt"
    log_path.write_text("\n".join(log_lines), encoding="utf-8")
    log(f"Saved execution log: {log_path.name}")

    receipt_path = EVIDENCE_DIR / "workbook_verification_receipt.json"
    receipt_path.write_text(json.dumps(verification_results, indent=2, ensure_ascii=False), encoding="utf-8")
    log(f"Saved workbook verification receipt: {receipt_path.name}")

if __name__ == "__main__":
    main()
